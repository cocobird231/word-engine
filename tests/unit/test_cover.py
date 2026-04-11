"""
Unit tests for the cover page renderer.

Tests cover: enabled/disabled toggle, field visibility flags, field values,
and page break insertion after cover.
"""
import pytest
from docx import Document
from src.renderer.cover import render_cover


def _doc_texts(doc):
    """Return all non-empty paragraph texts from a document."""
    return [p.text for p in doc.paragraphs if p.text.strip()]


def _minimal_params(**overrides):
    params = {
        "cover": {
            "enabled": True,
            "title_alignment": "center",
            "show_subtitle": True,
            "show_project_id": True,
            "show_version": True,
            "show_author": True,
            "show_organization": True,
            "show_date": True,
            "show_confidentiality": True,
        },
        "project": {
            "document_title": "Test Document",
            "subtitle": "A Test Subtitle",
            "project_id": "PROJ001",
            "version": "v1.0",
            "author": "蝦蝦",
            "organization": "蝦蝦隊",
            "updated_date": "2026-04-12",
            "confidentiality": "internal",
        },
    }
    # Apply any override keys
    for section, fields in overrides.items():
        if section in params:
            params[section].update(fields)
        else:
            params[section] = fields
    return params


class TestCoverEnabled:
    def test_cover_renders_title(self):
        doc = Document()
        render_cover(doc, _minimal_params())
        texts = _doc_texts(doc)
        assert "Test Document" in texts

    def test_cover_disabled_adds_nothing(self):
        doc = Document()
        params = _minimal_params()
        params["cover"]["enabled"] = False
        render_cover(doc, params)
        texts = _doc_texts(doc)
        assert texts == []

    def test_cover_adds_page_break(self):
        doc = Document()
        render_cover(doc, _minimal_params())
        # A page break paragraph exists (it may be empty text but has a run break)
        all_paragraphs = doc.paragraphs
        has_break = any(
            any("lastRenderedPageBreak" in r.element.xml or "w:br" in r.element.xml
                for r in p.runs)
            for p in all_paragraphs
        )
        assert has_break


class TestFieldVisibility:
    def test_subtitle_shown(self):
        doc = Document()
        render_cover(doc, _minimal_params())
        texts = _doc_texts(doc)
        assert "A Test Subtitle" in texts

    def test_subtitle_hidden(self):
        doc = Document()
        params = _minimal_params()
        params["cover"]["show_subtitle"] = False
        render_cover(doc, params)
        texts = _doc_texts(doc)
        assert "A Test Subtitle" not in texts

    def test_project_id_shown(self):
        doc = Document()
        render_cover(doc, _minimal_params())
        combined = " ".join(_doc_texts(doc))
        assert "PROJ001" in combined

    def test_project_id_hidden(self):
        doc = Document()
        params = _minimal_params()
        params["cover"]["show_project_id"] = False
        render_cover(doc, params)
        combined = " ".join(_doc_texts(doc))
        assert "PROJ001" not in combined

    def test_version_shown(self):
        doc = Document()
        render_cover(doc, _minimal_params())
        combined = " ".join(_doc_texts(doc))
        assert "v1.0" in combined

    def test_confidentiality_mapped_to_chinese(self):
        doc = Document()
        render_cover(doc, _minimal_params())
        combined = " ".join(_doc_texts(doc))
        assert "內部使用" in combined

    def test_date_shown(self):
        doc = Document()
        render_cover(doc, _minimal_params())
        combined = " ".join(_doc_texts(doc))
        assert "2026-04-12" in combined


class TestEdgeCases:
    def test_missing_cover_section_uses_defaults(self):
        doc = Document()
        params = {"project": {"document_title": "Minimal"}}
        render_cover(doc, params)
        texts = _doc_texts(doc)
        assert "Minimal" in texts

    def test_empty_document_title_shows_placeholder(self):
        doc = Document()
        params = _minimal_params()
        params["project"]["document_title"] = ""
        render_cover(doc, params)
        texts = _doc_texts(doc)
        assert any("No Title" in t for t in texts)
