"""
Unit tests for Phase 3 SEQ and REF field functionality.
Verifies OOXML structure of fields and integration with renderer.
"""
import pytest
from docx import Document
from docx.oxml.ns import qn
from src.renderer.seq_field import insert_seq_caption, insert_ref_field
from src.renderer.renderer import render_docx


def _get_instr_texts(paragraph):
    """Extract all instrText values from a paragraph."""
    return [e.text for e in paragraph._p.iter() if e.tag == qn("w:instrText")]


def _get_bookmarks(paragraph):
    """Extract all bookmark names from a paragraph."""
    return [e.get(qn("w:name")) for e in paragraph._p.iter()
            if e.tag == qn("w:bookmarkStart")]


def _params_with_crossref():
    return {
        "cover": {"enabled": False},
        "toc": {"enabled": False},
        "images": {
            "insert_caption": True, "caption_prefix": "圖",
            "numbering_mode": "flat", "chapter_separator": "-",
            "caption_position": "below", "max_width_cm": 14.0,
            "download_remote_images": False
        },
        "tables": {
            "caption_enabled": True, "caption_prefix": "表",
            "numbering_mode": "flat", "chapter_separator": "-",
            "caption_position": "above"
        },
        "cross_references": {
            "enabled": True, "style": "ieee", "hyperlink_enabled": True,
            "figure_reference_format": {"ieee": "圖 {number}"},
            "table_reference_format": {"ieee": "表 {number}"},
            "heading_reference_format": {"ieee": "第 {number} 節"},
            "bookmark_prefix_figure": "fig_",
            "bookmark_prefix_table": "tbl_",
            "bookmark_prefix_heading": "sec_",
        },
    }


class TestSeqFieldStructure:
    def test_insert_seq_creates_instrtext(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_seq_caption(p, "圖", "Test", "Figure", "1", 1, "fig_1")
        instr = _get_instr_texts(p)
        assert any("SEQ Figure" in (t or "") for t in instr)

    def test_insert_seq_includes_arabic_format(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_seq_caption(p, "圖", "Test", "Figure", "1", 1, "fig_1")
        instr = _get_instr_texts(p)
        assert any("ARABIC" in (t or "") for t in instr)

    def test_insert_seq_creates_bookmark(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_seq_caption(p, "圖", "Test Alt", "Figure", "1", 1, "fig_1")
        bookmarks = _get_bookmarks(p)
        assert "fig_1" in bookmarks

    def test_insert_seq_contains_prefix_and_alt_text(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_seq_caption(p, "圖", "系統架構", "Figure", "2", 2, "fig_2")
        assert "圖" in p.text
        assert "系統架構" in p.text

    def test_insert_seq_table(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_seq_caption(p, "表", "資料表", "Table", "1", 10, "tbl_1")
        instr = _get_instr_texts(p)
        assert any("SEQ Table" in (t or "") for t in instr)
        bookmarks = _get_bookmarks(p)
        assert "tbl_1" in bookmarks


class TestRefFieldStructure:
    def test_insert_ref_creates_instrtext(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_ref_field(p, "圖 1", "fig_1", with_hyperlink=True)
        instr = _get_instr_texts(p)
        assert any("REF fig_1" in (t or "") for t in instr)

    def test_insert_ref_has_hyperlink_flag(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_ref_field(p, "圖 1", "fig_1", with_hyperlink=True)
        instr = _get_instr_texts(p)
        assert any("\\h" in (t or "") for t in instr)

    def test_insert_ref_no_hyperlink(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_ref_field(p, "圖 1", "fig_1", with_hyperlink=False)
        # No hyperlink element
        has_hyperlink = any(e.tag == qn("w:hyperlink") for e in p._p.iter())
        assert not has_hyperlink

    def test_insert_ref_has_placeholder_text(self):
        doc = Document()
        p = doc.add_paragraph()
        insert_ref_field(p, "圖 5", "fig_5", with_hyperlink=False)
        assert "圖 5" in p.text


class TestIntegrationWithRenderer:
    def test_figure_caption_has_seq_field(self, tmp_path):
        out = str(tmp_path / "out.docx")
        md = "# Title\n\n## Ch1\n\n![Alt text](fake.png)\n"
        render_docx(md, _params_with_crossref(), out)
        doc = Document(out)
        for p in doc.paragraphs:
            instr = _get_instr_texts(p)
            if any("SEQ Figure" in (t or "") for t in instr):
                return  # Found SEQ Figure field
        pytest.fail("No SEQ Figure field found in output")

    def test_table_caption_has_seq_field(self, tmp_path):
        out = str(tmp_path / "out.docx")
        md = "# Title\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
        render_docx(md, _params_with_crossref(), out)
        doc = Document(out)
        for p in doc.paragraphs:
            instr = _get_instr_texts(p)
            if any("SEQ Table" in (t or "") for t in instr):
                return
        pytest.fail("No SEQ Table field found in output")

    def test_crossref_has_ref_field(self, tmp_path):
        out = str(tmp_path / "out.docx")
        md = "# Title\n\n如 {{ref:fig-1}} 所示。\n\n![Alt](fake.png)\n"
        render_docx(md, _params_with_crossref(), out)
        doc = Document(out)
        # Look for REF field in all paragraphs
        for p in doc.paragraphs:
            instr = _get_instr_texts(p)
            if any("REF" in (t or "") and "fig" in (t or "") for t in instr):
                return
        pytest.fail("No REF field found for cross-reference")

    def test_caption_bookmark_exists_for_ref_target(self, tmp_path):
        out = str(tmp_path / "out.docx")
        md = "# Title\n\n![Alt](fake.png)\n"
        render_docx(md, _params_with_crossref(), out)
        doc = Document(out)
        all_bookmarks = []
        for p in doc.paragraphs:
            all_bookmarks.extend(_get_bookmarks(p))
        # Should have a fig_ bookmark
        assert any("fig_" in (b or "") for b in all_bookmarks), \
            f"No fig_ bookmark found. Bookmarks: {all_bookmarks}"
