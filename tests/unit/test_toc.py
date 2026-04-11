"""
Unit tests for the TOC renderer.

The generated TOC is a Word-updatable field, not static text.
Tests verify: enabled/disabled, title presence, page breaks, and XML field structure.
"""
import pytest
from docx import Document
from docx.oxml.ns import qn
from src.renderer.toc import render_toc


def _minimal_toc_params(**overrides):
    params = {
        "toc": {
            "enabled": True,
            "title": "目錄",
            "levels": 3,
            "page_break_before": True,
            "use_hyperlinks": True,
        }
    }
    for key, val in overrides.items():
        params["toc"][key] = val
    return params


def _find_toc_field(doc):
    """Find the TOC field instrText element in the document XML."""
    for elem in doc.element.body.iter():
        if elem.tag == qn("w:instrText") and elem.text and "TOC" in elem.text:
            return elem
    return None


class TestTocEnabled:
    def test_toc_inserts_field(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params())
        assert _find_toc_field(doc) is not None

    def test_toc_disabled_inserts_nothing(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(enabled=False))
        assert _find_toc_field(doc) is None

    def test_toc_title_paragraph_exists(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params())
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert "目錄" in texts

    def test_toc_custom_title(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(title="Table of Contents"))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert "Table of Contents" in texts

    def test_toc_no_title_when_empty(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(title=""))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Should not crash and no title line
        assert "目錄" not in texts


class TestTocFieldContent:
    def test_toc_field_includes_level_range(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(levels=3))
        field = _find_toc_field(doc)
        assert field is not None
        assert '"1-3"' in field.text

    def test_toc_field_includes_hyperlink_flag(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(use_hyperlinks=True))
        field = _find_toc_field(doc)
        assert "\\h" in field.text

    def test_toc_field_no_hyperlink_flag(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(use_hyperlinks=False))
        field = _find_toc_field(doc)
        assert "\\h" not in field.text

    def test_toc_levels_clamped_max(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(levels=99))
        field = _find_toc_field(doc)
        assert '"1-6"' in field.text

    def test_toc_levels_clamped_min(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(levels=0))
        field = _find_toc_field(doc)
        assert '"1-1"' in field.text


class TestTocPageBreaks:
    def test_page_break_before(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(page_break_before=True))
        # Document should have at least one page break
        xml = doc.element.body.xml
        assert "w:lastRenderedPageBreak" in xml or "w:br" in xml

    def test_no_page_break_before(self):
        doc = Document()
        render_toc(doc, _minimal_toc_params(page_break_before=False))
        # Should still have TOC field
        assert _find_toc_field(doc) is not None
